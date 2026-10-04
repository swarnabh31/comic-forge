"""Turns a file the user attaches in the chat (PDF, DOCX, PNG, JPG, SVG) into
something the assistant can actually use:

- documents (PDF / DOCX)  -> extracted plain text the LLM can read as source material
- images (PNG / JPG)      -> a normalized RGB PNG, usable as a character reference
- SVG                     -> its text content (titles, labels) AND a rasterized PNG

Files are stored under config.ATTACHMENTS_DIR, which is intentionally not
mounted as a static directory: an uploaded SVG can contain scripts, so the
browser must never be able to open one from our origin. Only the normalized
PNG copies ever get promoted into the (served) comic folder.
"""
import re
import uuid
from io import BytesIO
from pathlib import Path

from PIL import Image

import config

IMAGE_EXTS = {".png", ".jpg", ".jpeg"}
ALLOWED_EXTS = IMAGE_EXTS | {".svg", ".pdf", ".docx"}
MAX_IMAGE_SIDE = 1024
MAX_IMAGE_PIXELS = 50_000_000     # decompression-bomb guard
MAX_PDF_PAGES = 40
TEXT_CAP = 20_000                  # per file, before the per-prompt budget is applied

FRIENDLY_TYPES = "PDF, DOCX, PNG, JPG or SVG"


class AttachmentError(ValueError):
    """Raised with a message that is safe to show directly to the user."""


def _looks_right(ext: str, data: bytes) -> bool:
    head = data[:8192]
    if ext == ".png":
        return head.startswith(b"\x89PNG\r\n\x1a\n")
    if ext in (".jpg", ".jpeg"):
        return head.startswith(b"\xff\xd8\xff")
    if ext == ".pdf":
        return b"%PDF-" in head[:1024]
    if ext == ".docx":
        return head.startswith(b"PK")
    if ext == ".svg":
        return re.search(rb"<svg", head, re.I) is not None
    return False


def _clean(text: str) -> str:
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _normalize_image(img: Image.Image) -> Image.Image:
    if img.mode in ("RGBA", "LA", "P"):
        img = img.convert("RGBA")
        bg = Image.new("RGBA", img.size, (255, 255, 255, 255))
        bg.alpha_composite(img)
        img = bg
    img = img.convert("RGB")
    if max(img.size) > MAX_IMAGE_SIDE:
        img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE), Image.LANCZOS)
    return img


def _load_raster(data: bytes, name: str) -> Image.Image:
    try:
        img = Image.open(BytesIO(data))
        if img.width * img.height > MAX_IMAGE_PIXELS:
            raise AttachmentError(f"'{name}' has too many pixels to process safely.")
        img.load()
    except AttachmentError:
        raise
    except Exception:
        raise AttachmentError(f"'{name}' couldn't be read as an image — the file may be corrupted.")
    return _normalize_image(img)


def _pymupdf():
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        try:
            import fitz as pymupdf  # older PyMuPDF releases
            return pymupdf
        except ImportError:
            raise AttachmentError("PDF and SVG support needs PyMuPDF on the server: pip install pymupdf")


def _svg_text(data: bytes) -> str:
    """Regex, not an XML parser, on purpose: no entity expansion to worry about."""
    src = data.decode("utf-8", errors="ignore")
    parts = re.findall(r"<(?:title|desc|text|tspan)[^>]*>(.*?)</(?:title|desc|text|tspan)>", src, re.S | re.I)
    text = "\n".join(re.sub(r"<[^>]+>", "", p).strip() for p in parts)
    return _clean(text)


def _rasterize_svg(data: bytes, name: str) -> Image.Image:
    mu = _pymupdf()
    try:
        doc = mu.open(stream=data, filetype="svg")
        page = doc[0]
        zoom = MAX_IMAGE_SIDE / max(page.rect.width, page.rect.height, 1)
        pix = page.get_pixmap(matrix=mu.Matrix(zoom, zoom), alpha=False)
        return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    except Exception:
        raise AttachmentError(f"'{name}' couldn't be rendered — is it a valid SVG?")


def _pdf_text(data: bytes, name: str) -> tuple[str, str]:
    mu = _pymupdf()
    try:
        doc = mu.open(stream=data, filetype="pdf")
    except Exception:
        raise AttachmentError(f"'{name}' couldn't be opened as a PDF.")
    if doc.needs_pass:
        raise AttachmentError(f"'{name}' is password-protected. Remove the password and attach it again.")
    pages = doc.page_count
    text = _clean("\n\n".join(doc[i].get_text() for i in range(min(pages, MAX_PDF_PAGES))))
    meta = f"{pages} page{'s' if pages != 1 else ''}" + (f", first {MAX_PDF_PAGES} read" if pages > MAX_PDF_PAGES else "")
    if not text:
        meta += ", no selectable text (scanned?)"
    return text, meta


def _docx_text(data: bytes, name: str) -> tuple[str, str]:
    try:
        import docx
    except ImportError:
        raise AttachmentError("Word support needs python-docx on the server: pip install python-docx")
    try:
        document = docx.Document(BytesIO(data))
    except Exception:
        raise AttachmentError(f"'{name}' couldn't be opened as a .docx file (older .doc files aren't supported — save as .docx).")
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    text = _clean("\n".join(lines))
    words = len(text.split())
    return text, f"{words} words" + ("" if text else ", no text found")


def process(filename: str, data: bytes, comic_id: str) -> dict:
    """Validate + parse one upload and store it. Returns a metadata dict (no id yet —
    comic_store assigns the per-comic id like "A1")."""
    name = Path(filename or "file").name[:120]
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise AttachmentError(f"'{name}' isn't a supported file type. Attach {FRIENDLY_TYPES}.")
    if not data:
        raise AttachmentError(f"'{name}' is empty.")
    if len(data) > config.MAX_ATTACHMENT_MB * 1024 * 1024:
        raise AttachmentError(f"'{name}' is larger than {config.MAX_ATTACHMENT_MB} MB.")
    if not _looks_right(ext, data):
        raise AttachmentError(f"'{name}' doesn't look like a real {ext[1:].upper()} file.")

    folder = config.ATTACHMENTS_DIR / comic_id
    folder.mkdir(parents=True, exist_ok=True)
    stem = uuid.uuid4().hex[:10]
    (folder / f"{stem}{ext}").write_bytes(data)

    text, meta, image = "", "", None
    if ext in IMAGE_EXTS:
        image = _load_raster(data, name)
        kind, meta = "image", f"{image.width}×{image.height}"
    elif ext == ".svg":
        text = _svg_text(data)
        image = _rasterize_svg(data, name)
        kind, meta = "svg", f"{image.width}×{image.height}"
    elif ext == ".pdf":
        text, meta = _pdf_text(data, name)
        kind = "pdf"
    else:
        text, meta = _docx_text(data, name)
        kind = "docx"

    image_file = None
    if image is not None:
        image_file = f"{stem}.png"
        image.save(folder / image_file)

    return {
        "name": name,
        "kind": kind,
        "meta": meta,
        "text": text[:TEXT_CAP],
        "truncated": len(text) > TEXT_CAP,
        "image_file": image_file,     # relative to ATTACHMENTS_DIR / comic_id
    }
